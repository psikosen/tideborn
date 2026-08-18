from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = "submission-docs/Tideborn-Design-Intent.docx"
BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
MUTED = RGBColor(0x55, 0x55, 0x55)


def set_font(run, size, color=RGBColor(0, 0, 0), bold=False, italic=False):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.right_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.1

for style_name, size, color, before, after in (
    ("Heading 1", 16, BLUE, 16, 8),
    ("Heading 2", 13, BLUE, 12, 6),
    ("Heading 3", 12, DARK_BLUE, 8, 4),
):
    style = doc.styles[style_name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(size)
    style.font.color.rgb = color
    style.font.bold = True
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

title = doc.add_paragraph()
title.paragraph_format.space_before = Pt(0)
title.paragraph_format.space_after = Pt(4)
set_font(title.add_run("Tideborn"), 25, DARK_BLUE, bold=True)

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_before = Pt(0)
subtitle.paragraph_format.space_after = Pt(16)
set_font(subtitle.add_run("Design Intent - Survival & Resource Management"), 12, MUTED)

sections = [
    (
        "Target Player",
        "Tideborn is for players who enjoy systemic survival games, expressive creature movement, environmental storytelling, and resource choices with visible consequences. The portrait prototype is designed for short mobile sessions but rewards curiosity: a new player can immediately swim, hunt, gather, and dig, while a systems-minded player can learn how shelter, food, tools, weather, predators, and ecosystems affect one another.",
    ),
    (
        "Game Concept",
        "The player is an intelligent future-evolved amphibious octopus on an ocean planet. Survival comes from behaving like an octopus rather than a human-shaped hero: swimming, jetting, gripping surfaces, camouflaging, hunting fish, manipulating resources with multiple arms, and excavating natural terrain into shelter. The central fantasy is to read a living environment and prepare a den before severe weather and winter arrive. Harvesting is physical and ecological. Removing a whole plant yields more material but weakens sediment; hunting changes local prey numbers; storms redistribute food, debris, animals, and danger.",
    ),
    (
        "Prototype Contents",
        "This single-player vertical slice contains a deterministic 2.5D coastal world rendered in Three.js. It includes granular diggable terrain, water and darkness shaders, terrain-occluded bioluminescent light, portrait touch controls, octopus swimming and tentacle flow, jet movement, grabbing, gripping, camouflage, ink, digging, fish hunting, minerals, tool switching, crafting, multiple claimable dens, den reinforcement, resource storage, day and night, changing seasons, ecosystem predation, breeding, creature lifespans, deep-sea habitats, and a major storm. A complete contest session escalates toward winter and ends with a clear result: establish a seven-den network or prepare one exceptional solo den, while avoiding death and catastrophic shelter failure.",
    ),
    (
        "Future Vision",
        "The full game expands the same architecture into a wrapped planetary belt with streamed continents, islands, reefs, caves, abyssal plains, and trenches. Regional simulation will keep climate, migration, food webs, erosion, volcanism, and player-built habitats alive off screen. Progression will emphasize biological adaptation and uniquely octopus technology: rope systems, underwater gardens, pressure adaptations, living light, tide-powered structures, and vast connected dens. A vulnerable coastal survivor can eventually become the ecological engineer of a world that remains dangerous, reactive, and far larger than the player.",
    ),
]

for heading, body in sections:
    doc.add_paragraph(heading, style="Heading 1")
    paragraph = doc.add_paragraph(body)
    paragraph.paragraph_format.widow_control = True

doc.core_properties.title = "Tideborn - Design Intent"
doc.core_properties.subject = "Survival & Resource Management prototype"
doc.core_properties.author = ""
doc.core_properties.last_modified_by = ""
doc.core_properties.keywords = ""
doc.save(OUTPUT)
